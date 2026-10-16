# -*- coding: utf-8 -*-
"""
Baut Unterrichtsmaterialien (Word .docx) aus JSON-Beschreibungen und schreibt das Manifest materialien.json.

  Eingabe:  <projekt>/<materialien.ordner>/*.json     (Standard: material/)
  Ausgabe:  <projekt>/<materialien.ausgabe>/<id>.docx + materialien.json   (Standard: formulare/)
  Aufruf:   python engine/werkzeuge/materialien_bauen.py <projekt> [--nur-pruefen]
  Benötigt: python-docx  (pip install python-docx)

Schema einer Materialdatei (Beispiel: engine/material/beispiel.json):
  id (optional, sonst Dateiname), titel, kurz, stufe, faecher[], zeit, thema, einheit, nach (data-title der Folie,
  hinter die eine Unterrichtsfolie gesetzt wird), lehrplanbezug (optional),
  feinziele[ „Die Schüler*innen können …“ ], klafki{gegenwart, zukunft, exemplarisch, struktur, zugaenglichkeit},
  ablauf[], aufgaben[{nr, text, items[], zeilen}], aufloesung[{nr, item, loesung, begruendung}],
  differenzierung (Text oder Liste), quellen[] (fertige APA-Angaben).
  Optional: bilder[{datei, breite_cm, unterschrift}] + bilder_spalten (unter dem Titel); je Aufgabe ebenso
  bilder/bilder_spalten (Raster ab 2 Spalten), neue_seite (true oder Zwischenüberschrift: Seitenumbruch davor),
  ohne_kaestchen (true: Aufzählung ohne ☐, z. B. für Lückentexte). Bilder brauchen Pillow.

Jedes Material besteht aus dem Arbeitsblatt und einer Seite „Für die Lehrperson“.
Prüfung: Feinziel-Form, Klafki-Pflichtfelder, vollständige Auflösung (jede Aufgabe, jede Zeile mit allen Feldern),
Platzhalter, gerade Anführungszeichen. Bei Problemen endet das Skript mit Fehlercode 1.
"""
import json, re, sys, os, glob
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
import konfig

KLAFKI = [('gegenwart', 'Gegenwartsbedeutung', 'Present significance'), ('zukunft', 'Zukunftsbedeutung', 'Future significance'),
          ('exemplarisch', 'Exemplarische Bedeutung', 'Exemplary significance'), ('struktur', 'Struktur des Inhalts', 'Structure of the content'),
          ('zugaenglichkeit', 'Zugänglichkeit und Darstellbarkeit', 'Accessibility')]
KLAFKI_APA = 'Klafki, W. (1958). Didaktische Analyse als Kern der Unterrichtsvorbereitung. Die Deutsche Schule, 50, 450–471.'
PFLICHT = ['titel', 'stufe', 'faecher', 'zeit', 'feinziele', 'klafki', 'ablauf', 'aufgaben', 'aufloesung']
VAGE = r'\b(verstehen|kennenlernen|kennen lernen|sensibilisier\w*|wissen|erfahren|erleben|bewusst werden|understand|know|appreciate|be aware)\b'
PLATZHALTER = [r'\bTODO\b', r'\bLorem\b', r'\bipsum\b', r'\bXXX\b', r'\bFIXME\b', r'\[Platzhalter']
TXT = {
    'de': dict(name='Name', datum='Datum', aufgabe='Aufgabe', lp='Für die Lehrperson', stufe='Schulstufe', faecher='Fächer',
               zeit='Zeitbedarf', lernziel='Lernziel', lehrplan='Lehrplanbezug', fz='Feinziele (operationalisiert)',
               klafki='Didaktische Analyse (nach Klafki)', ablauf='Ablauf', loes='Lösungshinweise und Erwartungshorizont',
               nr='Nr.', item='Item', loesung='Lösung', begr='Begründung', diff='Differenzierung', quellen='Quellen',
               praefix='Die Schüler*innen können '),
    'en': dict(name='Name', datum='Date', aufgabe='Task', lp='For the teacher', stufe='Level', faecher='Subjects',
               zeit='Time', lernziel='Objective', lehrplan='Curriculum', fz='Learning objectives (operationalised)',
               klafki='Didactic analysis (Klafki)', ablauf='Procedure', loes='Answers and expectations',
               nr='No.', item='Item', loesung='Answer', begr='Rationale', diff='Differentiation', quellen='References',
               praefix='Students can '),
}


# ------------------------------------------------------------------ Prüfung
def pruefe(m, praefix):
    """Prüft eine Materialbeschreibung; liefert eine Liste von Problemen (leer = in Ordnung)."""
    pr, mid = [], m.get('id', '?')
    for k in PFLICHT:
        if not m.get(k): pr.append(f'{mid}: Pflichtfeld „{k}“ fehlt oder ist leer')
    fz = m.get('feinziele') or []
    if fz and not 1 <= len(fz) <= 4: pr.append(f'{mid}: {len(fz)} Feinziele (erwartet 1 – 4)')
    for z in fz:
        if not z.startswith(praefix): pr.append(f'{mid}: Feinziel beginnt nicht mit „{praefix.strip()} …“: {z[:60]}')
        elif not re.search(r'\w', z[len(praefix):]): pr.append(f'{mid}: Feinziel ohne Inhalt')
        if re.search(VAGE, z.split('(')[0], re.I): pr.append(f'{mid}: vages Verb im Feinziel (nicht beobachtbar): {z[:60]}')
    kl = m.get('klafki') or {}
    for k, de, _ in KLAFKI:
        if not str(kl.get(k, '')).strip(): pr.append(f'{mid}: Klafki-Feld „{k}“ ({de}) fehlt')
    nrs = [a.get('nr') for a in m.get('aufgaben') or []]
    for a in m.get('aufgaben') or []:
        if a.get('nr') is None or not str(a.get('text', '')).strip(): pr.append(f'{mid}: Aufgabe ohne nr oder text: {a}')
    al = m.get('aufloesung') or []
    for z in al:
        for f in ('nr', 'item', 'loesung', 'begruendung'):
            if not str(z.get(f, '')).strip(): pr.append(f'{mid}: Auflösung zu Aufgabe {z.get("nr")}: Feld „{f}“ leer')
        if z.get('nr') not in nrs: pr.append(f'{mid}: Auflösung verweist auf unbekannte Aufgabe {z.get("nr")}')
    for n in nrs:
        if n not in [z.get('nr') for z in al]: pr.append(f'{mid}: Aufgabe {n} hat keine Auflösung')
        items = next((a.get('items') or [] for a in m['aufgaben'] if a.get('nr') == n), [])
        geloest = [z.get('item') for z in al if z.get('nr') == n]
        for it in items:
            if it not in geloest: pr.append(f'{mid}: Aufgabe {n}, Item „{it[:40]}“ ohne Auflösung')
    roh = json.dumps(m, ensure_ascii=False)
    for v in PLATZHALTER:
        if re.search(v, roh, re.I): pr.append(f'{mid}: Platzhalter {v}')
    if re.search(r'(?<!\\)\\"', roh): pr.append(f'{mid}: gerades Anführungszeichen im Text (typografische verwenden)')
    return pr


# ------------------------------------------------------------------ Word
def baue(m, ziel, K, tx):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    FARBE = RGBColor.from_string(K['materialien'].get('farbe', '1F5F5B'))
    bis = (lambda t: re.sub(r'(\d)\s*–\s*(\d)', r'\1 – \2', t)) if K['bis_leerzeichen'] else (lambda t: t)

    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
    st = d.styles['Normal']; st.font.name = K['materialien'].get('schrift', 'Arial'); st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(k), st.font.name)
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'de-DE' if K['sprache'] != 'en' else 'en-GB'); rpr.append(lang)

    def absatz(text='', size=None, bold=False, italic=False, color=None, after=4):
        p = d.add_paragraph()
        for i, t in enumerate(bis(str(text)).split('**')):   # **fett**
            if not t: continue
            r = p.add_run(t); r.bold = bold or i % 2 == 1; r.italic = italic
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)
        return p

    def ueberschrift(t, size=14):
        p = absatz(t, size=size, bold=True, color=FARBE, after=6); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
        return p

    def linie(p):
        # Unterkante und „between“: Word verschmilzt aufeinanderfolgende gleiche Rahmen sonst zu einem – Schreiblinien verschwänden
        pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr')
        for seite in ('w:bottom', 'w:between'):
            e = OxmlElement(seite)
            for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', '9A9A9A')): e.set(qn(k), v)
            b.append(e)
        pPr.append(b)

    def bildstrom(pfad):
        # python-docx kann kein WebP: jedes Bild als JPEG (max. 1600 px, Transparenz auf Weiß) einbetten
        from io import BytesIO
        from PIL import Image
        im = Image.open(pfad)
        if im.mode in ('RGBA', 'LA', 'P'):
            im = im.convert('RGBA'); weiss = Image.new('RGB', im.size, 'white'); weiss.paste(im, mask=im.split()[-1]); im = weiss
        im = im.convert('RGB'); im.thumbnail((1600, 1600))
        puffer = BytesIO(); im.save(puffer, 'JPEG', quality=88); puffer.seek(0)
        return puffer

    def bilder(liste, spalten=1):
        # liste: [{"datei": "img/…", "breite_cm": 12, "unterschrift": "…"}] – Pfade relativ zum Projekt; spalten > 1 = Raster
        da = []
        for b in liste or []:
            pfad = konfig.p(K, b['datei'])
            if os.path.exists(pfad): da.append((b, pfad))
            else: print('WARNUNG: Bild fehlt', b['datei'])
        if not da: return
        spalten = max(1, int(spalten))
        if spalten == 1:
            for b, pfad in da:
                d.add_picture(bildstrom(pfad), width=Cm(float(b.get('breite_cm', 12))))
                d.paragraphs[-1].paragraph_format.space_after = Pt(2)
                if b.get('unterschrift'): absatz(b['unterschrift'], size=8, italic=True, after=8)
            return
        raster = d.add_table(rows=(len(da) + spalten - 1) // spalten, cols=spalten)
        for i, (b, pfad) in enumerate(da):
            zelle = raster.rows[i // spalten].cells[i % spalten]
            zelle.paragraphs[0].add_run().add_picture(bildstrom(pfad), width=Cm(float(b.get('breite_cm', 16.5 / spalten - 0.5))))
            if b.get('unterschrift'):
                r = zelle.add_paragraph().add_run(bis(b['unterschrift'])); r.italic = True; r.font.size = Pt(8)
        d.add_paragraph().paragraph_format.space_after = Pt(2)

    def tabelle(zeilen, kopf=None, breiten=None):
        t = d.add_table(rows=0, cols=len(zeilen[0] if zeilen else kopf)); t.style = 'Table Grid'
        if kopf:
            c = t.add_row().cells
            for i, h in enumerate(kopf):
                c[i].text = ''; r = c[i].paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10)
        for z in zeilen:
            c = t.add_row().cells
            for i, v in enumerate(z):
                c[i].text = ''; r = c[i].paragraphs[0].add_run(bis(str(v))); r.font.size = Pt(10); r.bold = (i == 0 and not kopf)
        if breiten:
            for row in t.rows:
                for i, w in enumerate(breiten): row.cells[i].width = Cm(w)
        d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    # Fußzeile mit Titel
    fp = sec.footer.paragraphs[0]; fr = fp.add_run(m['titel']); fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)

    # --- Arbeitsblatt ---
    absatz(m['titel'], size=20, bold=True, color=FARBE, after=2)
    if m.get('kurz'): absatz(m['kurz'], size=11, italic=True, after=8)
    linie(absatz(f'{tx["name"]}: ______________________________    {tx["datum"]}: ______________', size=11, after=10))
    bilder(m.get('bilder'), m.get('bilder_spalten', 1))
    for a in m['aufgaben']:
        if a.get('neue_seite'):          # z. B. damit die Auflösung nicht schon beim Bearbeiten von Teil 1 sichtbar ist
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if isinstance(a['neue_seite'], str): absatz(a['neue_seite'], size=11, bold=True, color=FARBE, after=6)
        ueberschrift(f'{tx["aufgabe"]} {a["nr"]}', 13)
        absatz(a['text'], size=13)
        bilder(a.get('bilder'), a.get('bilder_spalten', 1))
        marke = '' if a.get('ohne_kaestchen') else '☐  '
        for it in a.get('items') or []:
            absatz(marke + it, size=13, after=2)
        for _ in range(int(a.get('zeilen', 0))):
            linie(absatz(' ', size=16, after=0))

    # --- Für die Lehrperson ---
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    absatz(tx['lp'] + ': ' + m['titel'], size=16, bold=True, color=FARBE, after=8)
    info = [(tx['stufe'], m['stufe']), (tx['faecher'], ', '.join(m['faecher'])), (tx['zeit'], m['zeit']), (tx['lernziel'], m['feinziele'][0])]
    if m.get('lehrplanbezug'): info.append((tx['lehrplan'], m['lehrplanbezug']))
    tabelle(info, breiten=[4, 13])
    ueberschrift(tx['fz'])
    for i, z in enumerate(m['feinziele'], 1): absatz(f'{i}. {z}', size=11, after=2)
    ueberschrift(tx['klafki'])
    for k, de, en in KLAFKI:
        absatz(f'**{de if K["sprache"] != "en" else en}:** {m["klafki"][k]}', size=11)
    ueberschrift(tx['ablauf'])
    for i, s in enumerate(m['ablauf'], 1): absatz(f'{i}. {s}', size=11, after=2)
    ueberschrift(tx['loes'])
    tabelle([(z['nr'], z['item'], z['loesung'], z['begruendung']) for z in m['aufloesung']],
            kopf=[tx['nr'], tx['item'], tx['loesung'], tx['begr']], breiten=[1.2, 4.5, 4.5, 6.8])
    if m.get('differenzierung'):
        ueberschrift(tx['diff'])
        for s in (m['differenzierung'] if isinstance(m['differenzierung'], list) else [m['differenzierung']]): absatz('– ' + s, size=11, after=2)
    ueberschrift(tx['quellen'])
    for q in list(m.get('quellen') or []) + [KLAFKI_APA]:
        p = d.add_paragraph(); r = p.add_run(q); r.font.size = Pt(10)   # APA-Seitenangaben ohne Bis-Leerzeichen
        p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
    d.save(ziel)
    return d


def pruefe_docx(pfad, m, tx):
    from docx import Document
    d = Document(pfad)
    t = [p.text for p in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells]
    roh = '\n'.join(t)
    return [f'{m["id"]}: „{s}“ fehlt im Dokument' for s in (tx['fz'], tx['klafki'], tx['loes'], 'Klafki, W. (1958)') if s not in roh]


def main(K, nur_pruefen=False):
    MK = K['materialien']
    tx = TXT.get(K['sprache'], TXT['de'])
    praefix = MK.get('feinziel_praefix') or tx['praefix']
    quelle = konfig.p(K, MK.get('ordner', 'material')); aus = konfig.p(K, MK.get('ausgabe', 'formulare'))
    dateien = sorted(glob.glob(os.path.join(quelle, '*.json')))
    if not dateien: print('Keine Materialbeschreibungen in', quelle); return 0
    os.makedirs(aus, exist_ok=True)
    manifest, probleme = [], []
    for f in dateien:
        m = json.load(open(f, encoding='utf-8'))
        m.setdefault('id', os.path.splitext(os.path.basename(f))[0])
        pr = pruefe(m, praefix); probleme += pr
        if nur_pruefen or any('Pflichtfeld' in x for x in pr): continue
        datei = re.sub(r'[^\w-]+', '_', m['id']) + '.docx'
        baue(m, os.path.join(aus, datei), K, tx)
        probleme += pruefe_docx(os.path.join(aus, datei), m, tx)
        manifest.append({'datei': datei, 'id': m['id'], 'titel': m['titel'], 'kurz': m.get('kurz', ''), 'einheit': m.get('einheit', ''),
                         'thema': m.get('thema', ''), 'nach': m.get('nach', ''), 'stufe': m['stufe'], 'faecher': m['faecher'],
                         'zeit': m['zeit'], 'feinziele': m['feinziele'], 'ablauf': m['ablauf'], 'lehrplanbezug': m.get('lehrplanbezug', '')})
    if not nur_pruefen:
        json.dump(manifest, open(os.path.join(aus, 'materialien.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
        print(f'{len(manifest)} Material(ien) + materialien.json -> {aus}')
    print('Prüfung:', 'keine Auffälligkeiten' if not probleme else f'{len(probleme)} Problem(e)')
    for p in probleme: print('  ', p)
    return 1 if probleme else 0


if __name__ == '__main__':
    sys.exit(main(konfig.lade(), '--nur-pruefen' in sys.argv))
