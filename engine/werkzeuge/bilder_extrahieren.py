# Bestandskatalog einer Vorlage: zieht aus PPTX-Dateien je Folie Text, Notizen, Links und Bilder.
# Ergebnis: <projekt>/img/orig/<kürzel>/sNN_M.webp und <projekt>/img/orig/katalog.json
# Der Katalog ist die Grundlage für den Inhaltsschalter „exakt nach Vorlage“: Nichts aus der Vorlage darf verloren gehen.
#
# Aufruf: python engine/werkzeuge/bilder_extrahieren.py <projekt> <datei.pptx | ordner-mit-pptx> [...]
# Benötigt: python-pptx, Pillow
import os, sys, io, json, re, glob
try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.exit('Bitte installieren: pip install python-pptx Pillow')
from PIL import Image

if len(sys.argv) < 3:
    sys.exit('Aufruf: python engine/werkzeuge/bilder_extrahieren.py <projekt> <datei.pptx | ordner> [...]')
PROJ = os.path.abspath(sys.argv[1])
dateien = []
for a in sys.argv[2:]:
    dateien += sorted(glob.glob(os.path.join(a, '*.pptx'))) if os.path.isdir(a) else [a]
if not dateien: sys.exit('Keine PPTX-Dateien gefunden.')

ZIEL = os.path.join(PROJ, 'img', 'orig'); os.makedirs(ZIEL, exist_ok=True)

def kuerzel(pfad, n):
    """Kurzes, dateisystemtaugliches Kürzel je Foliensatz (z. B. 'V1', 'V2' oder aus dem Dateinamen)."""
    stamm = re.sub(r'[^A-Za-z0-9]+', '', os.path.splitext(os.path.basename(pfad))[0])[:12]
    return stamm or f'V{n}'

def formen(shapes):
    """Alle Formen, auch in Gruppen."""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from formen(sh.shapes)
        else:
            yield sh

katalog = []
for n, pfad in enumerate(dateien, 1):
    kz = kuerzel(pfad, n); os.makedirs(os.path.join(ZIEL, kz), exist_ok=True)
    pr = Presentation(pfad)
    for i, sl in enumerate(pr.slides, 1):
        texte, links, bilder, m = [], [], [], 0
        for sh in formen(sl.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                texte.append(sh.text_frame.text.strip())
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.hyperlink is not None and r.hyperlink.address:
                            links.append({'text': r.text.strip(), 'url': r.hyperlink.address})
            try:
                if sh.click_action and sh.click_action.hyperlink and sh.click_action.hyperlink.address:
                    links.append({'text': getattr(sh, 'name', ''), 'url': sh.click_action.hyperlink.address})
            except Exception:
                pass
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob, ext = sh.image.blob, sh.image.ext
                except Exception:
                    continue
                m += 1; name = f's{i:02d}_{m}.webp'
                try:
                    im = Image.open(io.BytesIO(blob))
                    if im.mode in ('P', 'LA'): im = im.convert('RGBA')
                    if im.mode not in ('RGB', 'RGBA'): im = im.convert('RGB')
                    im.thumbnail((1600, 1600)); im.save(os.path.join(ZIEL, kz, name), 'WEBP', quality=82)
                    bilder.append({'datei': f'img/orig/{kz}/{name}', 'w': im.width, 'h': im.height, 'format': ext,
                                   'herkunft': '', 'lizenz': ''})   # Herkunft und Lizenz recherchieren und eintragen
                except Exception as e:
                    bilder.append({'datei': None, 'format': ext, 'hinweis': f'nicht lesbar ({ext}): {e}'})
        notizen = ''
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame is not None:
            notizen = sl.notes_slide.notes_text_frame.text.strip()
        katalog.append({'satz': kz, 'datei': os.path.basename(pfad), 'folie': i, 'text': '\n'.join(texte),
                        'notizen': notizen, 'links': links, 'bilder': bilder, 'umgesetzt_in': ''})
    print(f'{os.path.basename(pfad)}: {len(pr.slides)} Folien -> img/orig/{kz}/')

json.dump(katalog, open(os.path.join(ZIEL, 'katalog.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
nb = sum(len(k['bilder']) for k in katalog); nl = sum(len(k['links']) for k in katalog)
print(f'{len(katalog)} Folien, {nb} Bilder, {nl} Links -> img/orig/katalog.json')
print('Nächster Schritt: Für jedes Bild Herkunft und Lizenz eintragen, für jede Folie „umgesetzt_in“ (Ziel-Folie) pflegen.')
